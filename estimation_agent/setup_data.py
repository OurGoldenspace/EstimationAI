"""
setup_data.py  (v2)
===================
SOURCE 1 — Dashboard (454+ projects) → projects.json
SOURCE 2 — Individual workbooks      → divisions.json

Run from: estimation_agent/ folder (where all Excel files live)
Output:   data/projects.json
          data/divisions.json
"""

import os, json, re
from datetime import datetime
import openpyxl
from collections import Counter

EXCEL_DIR          = "."
OUTPUT_DIR         = "data"
DASHBOARD_FILENAME = "2026_Estimate_Summary_Dashboard_-_AI_Agent.xlsx"

DIVISION_CODES = {
    "01 00 00": "General Requirements",
    "02 00 00": "Existing Conditions",
    "03 00 00": "Concrete",
    "04 00 00": "Masonry",
    "05 00 00": "Metals",
    "06 00 00": "Wood, Plastics, and Composites",
    "07 00 00": "Thermal and Moisture Protection",
    "08 00 00": "Openings",
    "09 00 00": "Finishes",
    "10 00 00": "Specialties",
    "11 00 00": "Equipment",
    "12 00 00": "Furnishings",
    "13 00 00": "Special Construction",
    "14 00 00": "Conveying Equipment",
    "21 00 00": "Fire Suppression",
    "22 00 00": "Plumbing",
    "23 00 00": "HVAC",
    "25 00 00": "Integrated Automation",
    "26 00 00": "Electrical",
    "27 00 00": "Communications",
    "28 00 00": "Electronic Safety and Security",
    "31 00 00": "Earthwork",
    "32 00 00": "Exterior Improvements",
    "33 00 00": "Utilities",
}

# ─── HELPERS ───────────────────────────────────────────────────────────────────

def clean(v):
    if v is None: return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v

def flt(v):
    try: return float(v) if v is not None else None
    except: return None

def normalize_code(raw):
    if raw is None: return None
    return re.sub(r'\s+', ' ', str(raw).strip().rstrip())

def infer_type(name, raw_type):
    """Map raw type or name keywords to canonical project type."""
    if raw_type is not None:
        rt = str(raw_type).strip()
        mapping = {
            "Multi Residential": "Multi-Residential",
            "multi residential": "Multi-Residential",
            "Office":            "Office",
            "Retail":            "Retail",
            "Food":              "Grocery / Food Retail",
            "0":                 "Multi-Residential",
            "1":                 "Office",
            "2":                 "Retail",
            "3":                 "Institutional",
            "4":                 "Industrial",
            "5":                 "Mixed Use",
        }
        if rt in mapping:
            return mapping[rt]

    if not name:
        return "Commercial / Other"
    n = name.lower()

    if any(k in n for k in ["apartment","condo","residential","multi","ymca","daycare","school"]):
        return "Multi-Residential"
    if any(k in n for k in ["dental","medical","clinic","wellness","health","pharmacy","chiro"]):
        return "Medical / Dental"
    if any(k in n for k in ["sobeys","foodland","grocery","cannabis","anbl","dollarama","walmart","wal-mart"]):
        return "Grocery / Food Retail"
    if any(k in n for k in ["ford","honda","auto","dealership","toyota","nissan"]):
        return "Automotive"
    if any(k in n for k in ["reno","renovation","expansion","fit-up","tenant","facade","façade","exterior","upgrade"]):
        return "Renovation / Tenant Fit-up"
    if any(k in n for k in ["office","professional","stantec","bioscript"]):
        return "Office"
    if any(k in n for k in ["retail","store","shop","mall","bank","scotia"]):
        return "Retail"
    if any(k in n for k in ["warehouse","industrial","air liquide","manufacturing"]):
        return "Industrial"
    return "Commercial / Other"

def price_cat(val, price=None):
    if val:
        c = str(val).strip().upper()
        if c in ["A","B","C","D","E"]:
            return c
    if price:
        try:
            p = float(price)
            if p >= 5_000_000: return "A"
            if p >= 2_000_000: return "B"
            if p >= 500_000:   return "C"
            if p >= 100_000:   return "D"
            return "E"
        except: pass
    return None

def find_header_row(rows, marker="Closing Date"):
    """Find the row index where the header row starts."""
    for i, row in enumerate(rows):
        if any(v is not None and marker in str(v) for v in row):
            return i
    return None


# ══════════════════════════════════════════════════════════════════════════════
# SOURCE 1 — Dashboard → projects.json
# ══════════════════════════════════════════════════════════════════════════════

def read_dashboard(path):
    print(f"\n📊 Reading dashboard: {os.path.basename(path)}")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  ✗ Could not open: {e}")
        return []

    if "DATA" not in wb.sheetnames:
        print("  ✗ No DATA sheet found.")
        return []

    ws   = wb["DATA"]
    rows = list(ws.iter_rows(values_only=True))

    hdr_idx = find_header_row(rows)
    if hdr_idx is None:
        print("  ✗ Could not find header row.")
        return []

    hdr = rows[hdr_idx]
    col = {str(h).strip(): i for i, h in enumerate(hdr) if h}

    projects, skipped = [], 0
    for row in rows[hdr_idx + 1:]:
        if not any(v is not None for v in row):
            continue

        name = clean(row[col.get("Estimate Name", 6)])
        if not name:
            skipped += 1
            continue

        status = clean(row[col.get("Status", 17)])
        if status and status.lower() == "pending":
            skipped += 1
            continue

        raw_type  = clean(row[col.get("Type of Construction", 7)])
        price_raw = row[col.get("Estimate Price", 11)]
        area_raw  = row[col.get("Area", 8)]
        margin_raw= row[col.get("Margin", 16)]
        sched_raw = row[col.get("Schedule", 13)]
        sc_raw    = row[col.get("Soft Cost", 14)]
        cpw_raw   = row[col.get("Cost per Week", 15)]
        cd_raw    = row[col.get("Closing Date", 0)]

        # Try both "Estimate Number" and "Estimate Number2"
        est_num_raw = row[col.get("Estimate Number2", col.get("Estimate Number", 5))]

        price   = flt(price_raw)
        area_sf = flt(area_raw)
        margin  = flt(margin_raw)
        sc      = flt(sc_raw)
        cpw     = flt(cpw_raw)

        try: sched = int(float(sched_raw)) if sched_raw else None
        except: sched = None

        cpf    = round(price / area_sf, 2) if price and area_sf and area_sf > 0 else None
        cd_str = cd_raw.strftime("%Y-%m-%d") if isinstance(cd_raw, datetime) else None

        try: budget_year = int(row[col.get("Budget Year", 1)]) if row[col.get("Budget Year", 1)] else None
        except: budget_year = None

        projects.append({
            "source":          "dashboard",
            "estimate_number": str(clean(est_num_raw)) if est_num_raw else None,
            "estimate_name":   name,
            "project_type":    infer_type(name, raw_type),
            "raw_type":        raw_type,
            "area_sf":         area_sf,
            "province":        clean(row[col.get("Province", 3)]),
            "city":            clean(row[col.get("City/Town", 4)]),
            "budget_year":     budget_year,
            "quarter":         clean(row[col.get("Quarter", 2)]),
            "closing_date":    cd_str,
            "contract_type":   clean(row[col.get("Contract Type", 9)]),
            "payment_method":  clean(row[col.get("Payment Method", 10)]),
            "estimate_price":  round(price, 2) if price else None,
            "price_category":  price_cat(clean(row[col.get("Price Category", 12)]), price),
            "cost_per_sf":     cpf,
            "schedule_weeks":  sched,
            "soft_cost":       round(sc, 2) if sc else None,
            "cost_per_week":   round(cpw, 2) if cpw else None,
            "margin":          round(margin, 4) if margin else None,
            "margin_pct":      round(margin * 100, 2) if margin else None,
            "status":          status,
            "result":          clean(row[col.get("Successful / Unsucessful", 19)]),
        })

    print(f"  ✓ {len(projects)} projects  ({skipped} skipped)")
    return projects


# ══════════════════════════════════════════════════════════════════════════════
# SOURCE 2 — Individual workbooks → divisions.json
# ══════════════════════════════════════════════════════════════════════════════

def read_workbooks(excel_dir, dashboard_name):
    print(f"\n📁 Reading individual workbooks...")
    all_divs, processed = [], 0

    files = [
        f for f in os.listdir(excel_dir)
        if f.endswith(".xlsx")
        and not f.startswith("~")
        and f != dashboard_name
        and "Summary_Dashboard" not in f
    ]
    print(f"  Found {len(files)} workbooks")

    for fname in sorted(files):
        fpath = os.path.join(excel_dir, fname)
        try:
            wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        except Exception as e:
            print(f"  ✗ {fname}: {e}")
            continue

        if "Estimate Summary" not in wb.sheetnames or "DATA" not in wb.sheetnames:
            print(f"  ✗ {fname}: missing required sheets")
            continue

        # ── Read metadata from DATA sheet using header row ────────────────────
        est_num = est_name = province = city = raw_type = None
        area_sf = budget_year = price = None

        data_rows = list(wb["DATA"].iter_rows(values_only=True))
        hdr_idx = find_header_row(data_rows)

        if hdr_idx is not None and hdr_idx + 1 < len(data_rows):
            hdr = data_rows[hdr_idx]
            col = {str(h).strip(): i for i, h in enumerate(hdr) if h}
            row = data_rows[hdr_idx + 1]  # first data row after header
            try:
                # Handle both "Estimate Number" and "Estimate Number2"
                en_idx   = col.get("Estimate Number", col.get("Estimate Number2", 5))
                est_num  = str(row[en_idx]).strip() if row[en_idx] else None
                est_name = str(row[col.get("Estimate Name", 6)]).strip() if row[col.get("Estimate Name", 6)] else None
                raw_type = row[col.get("Type of Construction", 7)]
                area_sf  = float(row[col.get("Area", 8)]) if row[col.get("Area", 8)] else None
                province = str(row[col.get("Province", 3)]).strip() if row[col.get("Province", 3)] else None
                city     = str(row[col.get("City/Town", 4)]).strip() if row[col.get("City/Town", 4)] else None
                budget_year = int(row[col.get("Budget Year", 1)]) if row[col.get("Budget Year", 1)] else None
                price    = float(row[col.get("Estimate Price", 11)]) if row[col.get("Estimate Price", 11)] else None
            except Exception as e:
                print(f"    metadata parse error: {e}")

        if not est_name:
            print(f"  ✗ {fname}: no metadata found")
            continue

        project_type = infer_type(est_name, raw_type)

        # ── Read divisions from Estimate Summary ──────────────────────────────
        divs = []
        for row in wb["Estimate Summary"].iter_rows(values_only=True):
            if not row[0]:
                continue
            code = normalize_code(str(row[0]))
            if not re.match(r'^\d{2} 00 00$', code):
                continue
            if code not in DIVISION_CODES:
                continue

            try: total = float(row[7]) if row[7] else 0
            except: total = 0
            if total == 0:
                continue  # skip zero — means scope not included

            # Parse $/SF from col I (index 8)
            cpf = 0
            sf_raw = row[8] if len(row) > 8 else None
            if sf_raw:
                try: cpf = float(str(sf_raw).replace("$/SF", "").strip())
                except: pass
            if cpf == 0 and area_sf and area_sf > 0:
                cpf = round(total / area_sf, 2)

            def f(v):
                try: return float(v) if v else 0
                except: return 0

            divs.append({
                "estimate_number": est_num,
                "estimate_name":   est_name,
                "project_type":    project_type,
                "area_sf":         area_sf,
                "province":        province,
                "city":            city,
                "budget_year":     budget_year,
                "division_code":   code,
                "division_name":   DIVISION_CODES[code],
                "total":           round(total, 2),
                "cost_per_sf":     round(cpf, 2),
                "me":              round(f(row[3]), 2),
                "labour":          round(f(row[5]), 2),
                "subcontract":     round(f(row[6]), 2),
                "pct_of_project":  round(total / price * 100, 2) if price and price > 0 else 0,
            })

        if divs:
            all_divs.extend(divs)
            processed += 1
            print(f"  ✓ {est_name:<42} {len(divs):2d} divisions")
        else:
            print(f"  ✗ {fname}: no non-zero divisions found")

    print(f"\n  {processed} workbooks → {len(all_divs)} division records")
    return all_divs


# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    dashboard_path = os.path.join(EXCEL_DIR, DASHBOARD_FILENAME)
    if not os.path.exists(dashboard_path):
        print(f"⚠️  Dashboard not found: {dashboard_path}")
        print(f"   Place '{DASHBOARD_FILENAME}' in: {os.path.abspath(EXCEL_DIR)}")
        projects = []
    else:
        projects = read_dashboard(dashboard_path)

    divisions = read_workbooks(EXCEL_DIR, DASHBOARD_FILENAME)

    # Write output
    with open(os.path.join(OUTPUT_DIR, "projects.json"), "w", encoding="utf-8") as f:
        json.dump(projects, f, indent=2, ensure_ascii=False)
    with open(os.path.join(OUTPUT_DIR, "divisions.json"), "w", encoding="utf-8") as f:
        json.dump(divisions, f, indent=2, ensure_ascii=False)

    # Summary
    print("\n" + "="*60)
    print(f"✅ projects.json  → {len(projects)} projects")
    print(f"✅ divisions.json → {len(divisions)} division records")
    print("="*60)

    if projects:
        type_counts = Counter(p["project_type"] for p in projects)
        print("\nProject type breakdown:")
        for t, n in sorted(type_counts.items(), key=lambda x: -x[1]):
            print(f"  {t:<38} {n:>4}")

        prov = Counter(p["province"] for p in projects if p["province"])
        print(f"\nProvinces: {dict(prov)}")

        years = [p["budget_year"] for p in projects if p["budget_year"]]
        if years:
            print(f"Year range: {min(years)} – {max(years)}")

    if divisions:
        workbooks = len(set(d["estimate_name"] for d in divisions))
        print(f"\nDivision data from {workbooks} individual workbooks")


if __name__ == "__main__":
    main()