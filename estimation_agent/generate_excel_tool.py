# ══════════════════════════════════════════════════════════════════════════════
# TOOL 5 — generate_estimate_excel
# Uses ADK Artifact system for proper file downloads
# ══════════════════════════════════════════════════════════════════════════════

import io
import base64
from google.genai import types


async def generate_estimate_excel(
    project_type: str,
    area_sf: float,
    province: str,
    city: str,
    comparable_projects: list[dict],
    divisions: list[dict],
    construction_cost: float,
    ohp_amount: float,
    soft_cost: float,
    estimate_price: float,
    margin_pct: float,
    price_category: str,
    cost_per_sf: float,
    schedule_weeks: int,
    tool_context=None,
) -> dict:
    """
    Generates a professional Excel spreadsheet and saves it as an ADK Artifact
    for proper file download in the ADK web UI.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from datetime import datetime
    import os

    # ── Build the workbook ──────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Estimate"

    subheader_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    subheader_font = Font(bold=True, size=11)
    title_font = Font(bold=True, size=14, color="1F4E78")
    currency_format = "$#,##0.00"
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15

    row = 1
    ws[f'A{row}'] = f"AGCM ESTIMATE - {project_type}"
    ws[f'A{row}'].font = title_font
    row += 2

    ws[f'A{row}'] = "PROJECT DETAILS"
    ws[f'A{row}'].font = subheader_font
    row += 1
    for label, value in [
        ("Project Type:", project_type),
        ("Area:", f"{area_sf:,.0f} SF"),
        ("Location:", f"{city}, {province}"),
        ("Estimate Date:", datetime.now().strftime("%B %d, %Y")),
    ]:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = Font(bold=True, size=10)
        row += 1

    row += 1
    ws[f'A{row}'] = "COMPARABLE PROJECTS"
    ws[f'A{row}'].font = subheader_font
    row += 1
    for col, header in enumerate(["Project Name", "Area (SF)", "Cost/SF", "Margin %", "Similarity"], 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = subheader_fill
        cell.font = subheader_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    row += 1
    for proj in comparable_projects:
        ws.cell(row=row, column=1).value = proj.get("estimate_name", "")
        ws.cell(row=row, column=2).value = proj.get("area_sf", 0)
        ws.cell(row=row, column=3).value = proj.get("cost_per_sf", 0)
        ws.cell(row=row, column=4).value = proj.get("margin_pct", 0)
        ws.cell(row=row, column=5).value = proj.get("similarity_score", 0)
        ws.cell(row=row, column=2).number_format = "#,##0"
        ws.cell(row=row, column=3).number_format = currency_format
        ws.cell(row=row, column=4).number_format = "0.00%"
        for col in range(1, 6):
            ws.cell(row=row, column=col).border = thin_border
        row += 1

    row += 1
    ws[f'A{row}'] = "CSI DIVISION BREAKDOWN"
    ws[f'A{row}'].font = subheader_font
    row += 1
    for col, header in enumerate(["Division Code", "Division Name", "$/SF", "Estimated Total"], 1):
        cell = ws.cell(row=row, column=col)
        cell.value = header
        cell.fill = subheader_fill
        cell.font = subheader_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    row += 1
    for div in divisions:
        ws.cell(row=row, column=1).value = div.get("division_code", "")
        ws.cell(row=row, column=2).value = div.get("division_name", "")
        ws.cell(row=row, column=3).value = div.get("cost_per_sf", 0)
        ws.cell(row=row, column=4).value = div.get("estimated_total", 0)
        ws.cell(row=row, column=3).number_format = currency_format
        ws.cell(row=row, column=4).number_format = currency_format
        for col in range(1, 5):
            ws.cell(row=row, column=col).border = thin_border
        row += 1

    row += 1
    ws[f'A{row}'] = "ESTIMATE SUMMARY"
    ws[f'A{row}'].font = subheader_font
    row += 1
    for label, value, fmt in [
        ("Construction Cost:", construction_cost, currency_format),
        ("OH&P (5%):", ohp_amount, currency_format),
        ("Soft Costs (5%):", soft_cost, currency_format),
    ]:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'B{row}'].number_format = fmt
        ws[f'A{row}'].font = Font(bold=True, size=10)
        row += 1
    row += 1
    ws[f'A{row}'] = "ESTIMATE PRICE:"
    ws[f'A{row}'].font = Font(bold=True, size=11, color="FF0000")
    ws[f'B{row}'] = estimate_price
    ws[f'B{row}'].number_format = currency_format
    ws[f'B{row}'].font = Font(bold=True, size=11, color="FF0000")
    row += 2
    for label, value, fmt in [
        ("Price Category:", price_category, None),
        ("Margin:", margin_pct / 100, "0.0%"),
        ("Cost per SF:", cost_per_sf, currency_format),
        ("Estimated Schedule:", f"{schedule_weeks} weeks", None),
    ]:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        if fmt:
            ws[f'B{row}'].number_format = fmt
        ws[f'A{row}'].font = Font(bold=True, size=10)
        row += 1

    # ── Save to bytes buffer (for ADK artifact) ─────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    file_bytes = buffer.read()

    # ── Also save to disk (backup) ──────────────────────────────────────────
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = project_type.replace(' ', '_').replace('/', '_').replace('\\', '_')
    filename = f"AGCM_Estimate_{safe_name}_{timestamp}.xlsx"

    output_dir = os.path.join(os.path.dirname(__file__), "estimates")
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)
    with open(filepath, 'wb') as f:
        f.write(file_bytes)

    # ── Save as ADK Artifact for download button ────────────────────────────
    if tool_context is not None:
        try:
            artifact_part = types.Part(
                inline_data=types.Blob(
                    data=file_bytes,
                    mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            )
            version = await tool_context.save_artifact(
                filename=filename,
                artifact=artifact_part
            )
            return {
                "status": "success",
                "file_name": filename,
                "file_path": filepath,
                "artifact_version": version,
                "message": f"File '{filename}' (version {version}) is ready for download.",
                "download_ready": True,
            }
        except Exception as e:
            # Fallback: return file path if artifact fails
            return {
                "status": "success",
                "file_name": filename,
                "file_path": filepath,
                "message": f"Excel saved to: {filepath}",
                "download_ready": True,
                "note": f"Artifact save failed ({e}), file saved locally instead.",
            }
    else:
        return {
            "status": "success",
            "file_name": filename,
            "file_path": filepath,
            "message": f"Excel saved to: {filepath}",
            "download_ready": True,
        }